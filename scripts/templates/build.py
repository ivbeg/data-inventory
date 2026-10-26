#!/usr/bin/env python3
"""Собирает из CSV-шаблонов и схемы реестра рабочую таблицу workbook.xlsx,
документы Word и архив data-inventory-templates.zip.

Запуск из корня репозитория: python3 scripts/templates/build.py [static/templates]
Нужны openpyxl и python-docx.
"""
import csv, json, os, sys, zipfile
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

T = sys.argv[1] if len(sys.argv) > 1 else 'static/templates'
BOOK = 'https://datainventory.begtin.tech'
MAXROW = 1000

# ---------------------------------------------------------------- справочники
LISTS = {
    'access_level': ('Режим доступа', [
        ('open', 'открытые данные: весь набор файлом или через программный интерфейс без регистрации и платы'),
        ('unstructured', 'неструктурированные сведения: опубликовано для чтения человеком — страницы, PDF, поиск по записи'),
        ('restricted', 'с ограничениями: регистрация, плата, соглашение, запрос, доступ не для всех'),
        ('closed', 'закрыто: законного способа получить снаружи нет, и это подтверждено'),
        ('unknown', 'не установлено: проверить не удалось — запишите, что проверяли'),
    ]),
    'access_restriction': ('Причина ограничения', [
        ('personal_data', 'персональные данные, личная и семейная тайна'),
        ('state_secret', 'государственная тайна'),
        ('official', 'служебная информация ограниченного распространения'),
        ('confidential', 'коммерческая, банковская, налоговая, врачебная, статистическая и иные тайны по закону'),
        ('third_party', 'права третьих лиц: договор, лицензия, авторские права'),
        ('owner_decision', 'решение владельца, нормы нет'),
        ('none', 'ограничений нет'),
        ('unknown', 'причина не установлена'),
    ]),
    'update_frequency': ('Периодичность обновления', [
        ('continuous', 'непрерывно'), ('daily', 'ежедневно'), ('weekly', 'еженедельно'),
        ('monthly', 'ежемесячно'), ('quarterly', 'ежеквартально'), ('annual', 'ежегодно'),
        ('irregular', 'нерегулярно'), ('never', 'не обновляется'), ('unknown', 'не установлено'),
    ]),
    'record_type': ('Тип записи', [
        ('dataset', 'набор данных — основная запись'), ('resource', 'ресурс системы'),
        ('system', 'информационная система — временная замена, пока не найдены наборы'),
        ('site', 'сайт — временная замена, пока не найдена система'),
        ('group', 'группа однотипных источников одной строкой'),
    ]),
    'yes_no_unknown': ('Да, нет, не знаю', [
        ('yes', 'да'), ('no', 'нет'), ('unknown', 'не знаю, не установлено'),
    ]),
    'record_status': ('Статус записи', [
        ('active', 'действует'), ('retired', 'выведена из эксплуатации, не ведётся'),
        ('url_broken', 'адрес не работает'),
    ]),
    'layer': ('Слой данных', [
        ('government', 'государственные'), ('public', 'общественные'), ('commercial', 'коммерческие'),
    ]),
    'owner_status': ('Статус владельца', [
        ('active', 'действует'), ('abolished', 'упразднён — укажите преемника'),
    ]),
    'mandate_result': ('Итог по полномочию', [
        ('found', 'найдено'), ('partial', 'найдено частично'), ('not_found', 'не найдено'),
        ('transferred', 'передано другому владельцу'),
    ]),
    'row_status': ('Отметка в анкете', [
        ('confirmed', 'подтверждаю: строка верна'), ('corrected', 'исправляю: строка верна с поправками'),
        ('new', 'новое: строку добавил владелец'), ('not_ours', 'не наше: набор ведёт другой'),
        ('not_kept', 'не ведём: набор не ведут или перестали вести'),
        ('needed', 'не хватает: данные нужны, но их нет'),
    ]),
    'storage': ('Где хранятся данные', [
        ('system', 'в информационной системе'), ('spreadsheet', 'в электронной таблице'),
        ('document', 'в документе'), ('email', 'в почте'), ('paper', 'на бумаге'),
    ]),
    'respondent_result': ('Итог опроса адресата', [
        ('answered', 'ответил'), ('no_data', 'ответил: данных не ведём'), ('refused', 'отказался'),
        ('no_answer', 'не ответил'), ('redirected', 'переадресовал — укажите кому'),
    ]),
    'dimension': ('Признак качества', [
        ('completeness', 'полнота'), ('uniqueness', 'уникальность'), ('timeliness', 'актуальность'),
        ('validity', 'валидность: значения по правилам и справочникам'), ('consistency', 'согласованность'),
        ('accuracy', 'точность'), ('format', 'пригодность формата'), ('fitness', 'годность для вопроса'),
    ]),
    'impact': ('Влияние находки', [
        ('none', 'не мешает'), ('caveat', 'оговорка'), ('blocking', 'мешает'),
    ]),
    'fitness': ('Годится ли для вопроса', [
        ('fit', 'годится'), ('fit_with_caveats', 'годится с оговорками'), ('unfit', 'не годится'),
        ('not_checked', 'не проверено — с причиной'),
    ]),
    'search_result': ('Итог поиска', [
        ('original', 'найден первоисточник'), ('copy', 'только копии'),
        ('lead', 'след: система, владелец, каталог'), ('none', 'ничего'),
    ]),
    'change': ('Вид изменения', [
        ('new', 'новая запись'), ('changed', 'изменено поле'), ('status', 'изменён статус'),
    ]),
    'theme': ('Типовая тема', [
        ('статистика', ''), ('организации', ''), ('финансы', ''), ('люди', ''), ('ресурсы', ''),
    ]),
    'gap_type': ('Вид пробела', [
        ('no_data', 'данных нет ни у кого'), ('closed', 'есть, но закрыты'),
        ('unfit', 'есть, но не годятся: неполны, устарели, нет нужного разреза'),
        ('no_owner', 'есть, но владелец не установлен'), ('unknown', 'не установлено: проверить не удалось'),
    ]),
    'decision_type': ('Решение', [
        ('request', 'запрос владельцу'), ('agreement', 'договор о доступе'),
        ('own_collection', 'собственный сбор'), ('system_change', 'доработка системы'),
        ('accept', 'принять до следующего обновления'),
    ]),
    'plan_status': ('Статус строки плана', [
        ('planned', 'запланировано'), ('in_progress', 'в работе'), ('done', 'сделано'), ('cancelled', 'отменено'),
    ]),
    'priority': ('Приоритет вопроса', [
        ('главный', 'главный вопрос цели'), ('попутный', 'попутный вопрос'),
    ]),
}

# ---------------------------------------------------------------- листы
# (код, название, подсказка, обязательно, справочник или ссылка)
# обязательно: номер атрибута 1–14, 'всегда' или ''
# справочник: 'list:имя' (строго), 'list?:имя' (с предупреждением), 'ref:Лист' (ссылка на id листа)
REGISTRY = [
    ('id', 'Идентификатор', 'Постоянный код латиницей, цифрами и дефисами: faa-aircraft-registry. Не меняется при переименовании и переезде; не переиспользуется.', 1, None),
    ('title', 'Наименование', 'Официальное название набора или ресурса — из акта, страницы набора.', 2, None),
    ('description', 'Описание', 'Что описывает одна строка набора, за какой период и территорию, записи это или сводные показатели, чего в наборе нет.', 3, None),
    ('owner_person', 'Ответственный', 'Сотрудник, который знает данные: профильное подразделение или ИТ, не пресс-служба. Если не нашли — «не установлено» и где искать.', 4, None),
    ('owner_contacts', 'Контакты ответственного', 'Рабочая почта и телефон ответственного, а не приёмная.', 5, None),
    ('update_frequency', 'Периодичность обновления', 'Код из списка. Проверьте по датам в самих данных.', 6, 'list:update_frequency'),
    ('access_level', 'Режим доступа', 'Код шкалы доступа — по самому доступному способу получить данные, который проверили на деле.', 7, 'list:access_level'),
    ('access_basis', 'Чем подтверждён доступ', 'Ссылка, акт с пунктом, ответ владельца с датой; у unknown — что проверяли, из какой сети и что дальше.', '', None),
    ('access_restriction', 'Причина ограничения', 'Код из списка; несколько — через «;». У restricted и closed обязательно.', '', 'list?:access_restriction'),
    ('access_conditions', 'Условия доступа', 'У restricted: кто, как, за сколько и за какой срок получает данные и что выдают.', '', None),
    ('access_checked', 'Дата проверки доступа', 'ГГГГ-ММ-ДД или ГГГГ-ММ.', '', None),
    ('download_url', 'Ссылка для выгрузки', 'Прямой адрес файла или программного интерфейса; если набор не опубликован — «нет — набор не опубликован».', 8, None),
    ('format', 'Формат распространения', 'CSV, JSON, XML, XLSX, PDF, API; у ZIP — что внутри. По самому файлу, а не по подписи на странице.', 9, None),
    ('issued', 'Дата публикации', 'ГГГГ-ММ-ДД, «нет» или «не установлено»; пояснение — после « — ».', 10, None),
    ('modified', 'Дата изменения данных', 'Когда обновлены сами данные: ГГГГ-ММ-ДД, «нет» или «не установлено»; пояснение — после « — ».', 11, None),
    ('quality_check', 'Отметка о контроле качества', 'Когда и кто проверял, на чём, что найдено — числами, вопросы владельцу. Если не проверяли — «не проверено» и причина.', 12, None),
    ('license', 'Условия использования', 'Лицензия со ссылкой или условия соглашения; если их нет — «не указаны».', 13, None),
    ('publisher', 'Организация-владелец', 'Владелец по документу — id из листа «Владельцы». Колонка названа как в DCAT, но издатель на портале — не всегда владелец.', 14, 'ref?:Владельцы'),
    ('domain', 'Домен', 'Если реестр ведут домены: id строки подразделения из листа «Владельцы». Пусто, если реестр ведёт одна команда.', '', 'ref?:Владельцы'),
    ('aliases', 'Синонимы', 'Другие названия и прежние адреса, через «;».', '', None),
    ('type', 'Тип записи', 'Основная запись — набор; система, сайт и группа — временная замена.', '', 'list:record_type'),
    ('schema', 'Схема данных', 'Ссылка на описание полей или «нет».', '', None),
    ('data_contract', 'Контракт данных', 'Ссылка на контракт данных или соглашение об уровне обслуживания; «нет», если его нет.', '', None),
    ('information_system', 'Информационная система', 'Система, в которой ведут набор.', '', None),
    ('operator', 'Оператор системы', 'Кто эксплуатирует систему: учреждение, центр информатизации, подрядчик.', '', None),
    ('mandate', 'Полномочие', 'id строки из листа «Полномочия» или формулировка.', '', 'ref?:Полномочия'),
    ('legal_act', 'Нормативный акт', 'Закон, постановление, приказ — с номером, датой и пунктом.', '', None),
    ('personal_data', 'Персональные данные', 'Есть ли в наборе сведения о людях.', '', 'list:yes_no_unknown'),
    ('related', 'Связанные записи', 'aggregate_of:id, copy_of:id или successor_of:id; несколько — через «;».', '', None),
    ('status', 'Статус записи', 'Исчезнувшую запись не удаляют, а меняют статус и дают копию в веб-архиве.', '', 'list:record_status'),
    ('archive_url', 'Копия в веб-архиве', 'Адрес копии с датой; у исчезнувшей записи — последняя копия, где источник ещё работал.', '', None),
    ('known_from', 'Откуда известно', 'Метод и документ: «анализ полномочий; положение о министерстве». Подтверждение владельца — сюда же: «подтверждено владельцем: кто, чем, дата». Если запись составила языковая модель — какая, когда, подтверждена ли.', 'всегда', None),
    ('checked', 'Дата проверки записи', 'Когда составитель проверил запись: ГГГГ-ММ-ДД или ГГГГ-ММ.', 'всегда', None),
]

SHEETS = [
    # лист, файл, что в строке, глава (путь), колонки
    ('Реестр', 'registry.csv', 'набор данных или временная замена: ресурс, система, сайт, группа',
     '/process/registry/', REGISTRY),
    ('Владельцы', 'owners.csv', 'организация или подразделение', '/process/owners-and-mandates/', [
        ('id', 'Идентификатор', 'Код латиницей: faa, sfr-ru. На него ссылаются реестр, полномочия, адресаты.', 'всегда', None),
        ('name', 'Полное название', 'Официальное название организации или подразделения.', 'всегда', None),
        ('short_name', 'Короткое название', 'Сокращение: FAA, ФСС России.', '', None),
        ('level', 'Уровень', 'Центральный, региональный, муниципальный, подразделение, фонд, общественный проект.', '', None),
        ('layer', 'Слой данных', 'По основанию сбора и тому, у кого копия, а не по форме собственности.', '', 'list:layer'),
        ('parent', 'Вышестоящая организация', 'Название или id вышестоящей организации.', '', None),
        ('unit', 'Подразделение', 'Подразделение, если строка — подразделение организации.', '', None),
        ('contacts', 'Контакты', 'Адрес, телефон, почта для вопросов о данных.', '', None),
        ('mandate_document', 'Документ о полномочиях', 'Положение об органе, закон, устав.', '', None),
        ('status', 'Статус', 'Упразднённого владельца не удаляют: статус abolished и преемник.', '', 'list:owner_status'),
        ('successor', 'Преемник', 'id преемника из этого же листа.', '', 'ref?:Владельцы'),
        ('checked', 'Дата проверки', 'ГГГГ-ММ-ДД или ГГГГ-ММ.', 'всегда', None),
    ]),
    ('Полномочия', 'mandates.csv', 'полномочие владельца с цитатой из акта', '/process/owners-and-mandates/', [
        ('id', 'Идентификатор', 'Код полномочия латиницей; его пишут в колонку mandate реестра.', 'всегда', None),
        ('owner', 'Владелец', 'id из листа «Владельцы».', 'всегда', 'ref?:Владельцы'),
        ('mandate', 'Полномочие', 'Коротко: «регистрация воздушных судов».', 'всегда', None),
        ('quote', 'Цитата', 'Дословная цитата из акта — по ней видно, какие данные должны быть.', 'всегда', None),
        ('legal_act', 'Акт', 'Акт с номером и датой.', 'всегда', None),
        ('clause', 'Пункт', 'Статья, пункт, подпункт.', '', None),
        ('expected_data', 'Ожидаемые данные', 'Какие данные следуют из полномочия.', '', None),
        ('result', 'Итог', 'Нашли ли данные по этому полномочию.', '', 'list:mandate_result'),
        ('found', 'Что найдено', 'id найденных записей реестра, через «;».', '', None),
        ('known_from', 'Откуда известно', 'Метод и документ.', 'всегда', None),
        ('checked', 'Дата проверки', 'ГГГГ-ММ-ДД или ГГГГ-ММ.', 'всегда', None),
    ]),
    ('Адресаты', 'respondents.csv', 'адресат опроса владельцев', '/process/questionnaires/', [
        ('id', 'Идентификатор', 'Код адресата латиницей; его пишут в колонку respondent анкеты.', 'всегда', None),
        ('owner', 'Владелец', 'id организации из листа «Владельцы».', 'всегда', 'ref?:Владельцы'),
        ('unit', 'Подразделение', 'Профильное подразделение, ИТ, аналитика, защита данных, архив, закупки.', '', None),
        ('role', 'Роль', 'Зачем опрашиваем: профильное подразделение, оператор системы, подведомственная организация.', '', None),
        ('respondent', 'Адресат', 'Должность или ФИО.', '', None),
        ('contacts', 'Контакты', 'Почта и телефон.', '', None),
        ('sent', 'Отправлено', 'Дата отправки анкеты, ГГГГ-ММ-ДД.', '', None),
        ('reminded', 'Напоминание', 'Дата напоминания — до срока ответа.', '', None),
        ('answered', 'Ответ', 'Дата ответа.', '', None),
        ('interview', 'Интервью', 'Дата интервью — назначают при рассылке.', '', None),
        ('result', 'Итог', 'Ответ «данных не ведём» тоже записывают.', '', 'list:respondent_result'),
        ('redirected_to', 'Кому переадресовано', 'Новый адресат — id из этого листа или название.', '', None),
        ('checked', 'Дата проверки', 'ГГГГ-ММ-ДД или ГГГГ-ММ.', 'всегда', None),
    ]),
    ('Анкета', 'questionnaire.csv', 'ответ владельца о наборе', '/process/questionnaires/', [
        ('id', 'Идентификатор', 'У строки, вписанной заранее, — id записи реестра; у новой — пусто, присвоит составитель.', '', None),
        ('respondent', 'Адресат', 'id из листа «Адресаты».', 'всегда', 'ref?:Адресаты'),
        ('row_status', 'Отметка', 'Обязательна у каждой строки: строка без отметки — не ответ.', 'всегда', 'list:row_status'),
        ('prefilled_from', 'Откуда строка', 'У вписанных заранее: полномочие, контракт, адрес страницы.', '', None),
        ('system', 'Система', 'Вопрос 1: в какой системе ведётся набор.', '', None),
        ('storage', 'Где хранится', 'Вопрос 1: если вне систем — таблица, документ, почта, бумага.', '', 'list:storage'),
        ('dataset', 'Набор', 'Вопрос 2: как называется набор.', '', None),
        ('description', 'Что в наборе', 'Вопрос 2: люди, организации, объекты или события; период и территория.', '', None),
        ('owner_person', 'Кто ведёт', 'Вопрос 3: ФИО и должность.', '', None),
        ('owner_contacts', 'Контакты', 'Вопрос 3: почта и телефон.', '', None),
        ('update_frequency', 'Как часто обновляется', 'Вопрос 4.', '', 'list:update_frequency'),
        ('received_from', 'Откуда поступают данные', 'Вопрос 5: вносят сами, получают от других, покупают.', '', None),
        ('shared_with', 'Кому передаются', 'Вопрос 6: другим ведомствам, гражданам, публично; есть ли ссылка.', '', None),
        ('personal_data', 'Сведения о людях', 'Вопрос 7.', '', 'list:yes_no_unknown'),
        ('access_restrictions', 'Ограничения доступа', 'Вопрос 8: тайна, договор и т. п.', '', None),
        ('restriction_basis', 'Основание ограничения', 'Вопрос 8: норма или договор.', '', None),
        ('schema_available', 'Есть ли описание структуры', 'Вопрос 9: схема, справочник полей.', '', 'list:yes_no_unknown'),
        ('comment', 'Комментарий', 'У «не наше» — у кого набор; у «не ведём» — когда перестали и куда делись данные.', '', None),
        ('answered', 'Дата ответа', 'ГГГГ-ММ-ДД.', '', None),
    ]),
    ('Проверки качества', 'quality-checks.csv', 'одна проверка одного набора', '/process/quality/', [
        ('id', 'Идентификатор', 'Код проверки.', 'всегда', None),
        ('record', 'Запись', 'id записи реестра.', 'всегда', 'ref?:Реестр'),
        ('question', 'Вопрос', 'Номер вопроса, если проверка нужна только ему.', '', 'ref?:Вопросы'),
        ('dimension', 'Признак', 'Признак качества из списка.', 'всегда', 'list:dimension'),
        ('check', 'Что проверяли', '«Доля пустых значений в поле YEAR MFR».', 'всегда', None),
        ('threshold', 'Порог', 'Что приемлемо — записан до проверки.', '', None),
        ('result', 'Результат', 'Число или короткая находка.', 'всегда', None),
        ('impact', 'Влияние', 'Для вопроса, если он указан, иначе для набора в целом.', '', 'list:impact'),
        ('data_as_of', 'Данные на дату', 'Дата выгрузки, на которой проверяли.', '', None),
        ('checked_by', 'Кто проверял', 'Роль или имя.', '', None),
        ('checked', 'Дата проверки', 'ГГГГ-ММ-ДД.', 'всегда', None),
    ]),
    ('Вопросы', 'questions.csv', 'вопрос к данным', '/process/goals-and-questions/', [
        ('n', 'Номер', 'Постоянный номер: вопросы не перенумеровывают, снятый помечают.', 'всегда', None),
        ('text', 'Вопрос', 'В укрупнённой форме: «Где найти сведения о…?».', 'всегда', None),
        ('theme', 'Типовая тема', 'Статистика, организации, финансы, люди, ресурсы.', '', 'list?:theme'),
        ('task', 'Задача', 'Зачем нужен вопрос.', '', None),
        ('priority', 'Приоритет', 'Главный вопрос или попутный.', '', 'list?:priority'),
        ('threshold', 'Порог «годится»', 'Когда данных для вопроса достаточно — числами, до проверки. У главных вопросов обязателен, у попутных хватает общего порога.', '', None),
    ]),
    ('Источники', 'sources.csv', 'источник карты данных', '/sources/three-layers/', [
        ('id', 'Идентификатор', 'У сайта — домен с дефисами: med-workers-yearbook.', 'всегда', None),
        ('name', 'Название', 'Название или адрес источника.', 'всегда', None),
        ('type', 'Тип', 'Набор, ресурс, система, сайт или группа.', '', 'list?:record_type'),
        ('owner', 'Владелец', 'Организация, которая ведёт источник.', '', None),
        ('legal_basis', 'Основание', 'Закон или полномочие, по которому собирают данные.', '', None),
        ('level', 'Уровень', 'Федеральный, региональный, общественный проект.', '', None),
        ('layer', 'Слой данных', 'Государственные, общественные, коммерческие.', '', 'list:layer'),
        ('url', 'Адрес', 'Адрес источника.', '', None),
        ('aliases', 'Синонимы', 'Другие названия, через «;».', '', None),
        ('copy_of', 'Первоисточник', 'У копии — id первоисточника из этого листа: копию и первоисточник считают одним источником.', '', 'ref?:Источники'),
        ('known_from', 'Откуда известно', 'Метод и документ.', 'всегда', None),
    ]),
    ('Связи', 'links.csv', 'пара «вопрос — источник» или «вопрос — набор»', '/process/quality/', [
        ('question_n', 'Вопрос', 'Номер вопроса из листа «Вопросы».', 'всегда', 'ref?:Вопросы'),
        ('source_id', 'Источник', 'id из листа «Источники» или записи реестра.', 'всегда', 'ref?:Источники'),
        ('status', 'Доступ', 'Код шкалы доступа — для этого вопроса.', 'всегда', 'list:access_level'),
        ('status_basis', 'Чем подтверждён доступ', 'Ссылка, акт, ответ владельца с датой.', '', None),
        ('quality', 'Годится ли', 'Итог для пары по порогу вопроса.', '', 'list:fitness'),
        ('quality_note', 'Оговорка', 'Почему «с оговорками» или «не годится».', '', None),
        ('checked', 'Дата проверки', 'ГГГГ-ММ-ДД или ГГГГ-ММ.', 'всегда', None),
    ]),
    ('Журнал поиска', 'searches.csv', 'один запрос к поисковику или каталогу', '/sources/dataset-search/', [
        ('id', 'Идентификатор', 'Код запроса.', 'всегда', None),
        ('question_n', 'Вопрос', 'Номер вопроса; пусто, если поиск не под вопрос.', '', 'ref?:Вопросы'),
        ('engine', 'Поисковик или каталог', 'dateno, google-dataset-search, openaire, портал.', 'всегда', None),
        ('query', 'Запрос', 'Дословно.', '', None),
        ('lang', 'Язык', 'ru, en, kk…', '', None),
        ('filters', 'Фильтры', 'Страна, тип каталога, формат.', '', None),
        ('date', 'Дата', 'ГГГГ-ММ-ДД.', 'всегда', None),
        ('viewed', 'Просмотрено', 'Сколько результатов просмотрено.', '', None),
        ('result', 'Итог', 'Пустой результат тоже записывают.', 'всегда', 'list:search_result'),
        ('found', 'Что найдено', 'id записей реестра или источников, через «;».', '', None),
        ('note', 'Пометка', 'Что нашли и что делать дальше.', '', None),
    ]),
    ('План', 'plan.csv', 'пробел и решение, как его закрыть', '/process/gaps-and-plan/', [
        ('id', 'Идентификатор', 'Код строки плана.', 'всегда', None),
        ('question_n', 'Вопрос', 'Номер вопроса, которому мешает пробел.', '', 'ref?:Вопросы'),
        ('gap', 'Пробел', 'Чего не хватает — из таблицы покрытия.', 'всегда', None),
        ('gap_type', 'Вид пробела', 'Вид пробела решает, какое решение искать первым.', 'всегда', 'list:gap_type'),
        ('decision_type', 'Решение', 'Запрос владельцу, договор о доступе, собственный сбор, доработка системы или «принять».', 'всегда', 'list:decision_type'),
        ('decision', 'Что сделать', 'Что именно запросить, о чём договориться, что доработать.', '', None),
        ('who', 'Кто', 'Подразделение владельца или сотрудник по имени, а не «команда инвентаризации».', 'всегда', None),
        ('due', 'Срок', 'Дата или месяц: ГГГГ-ММ-ДД, ГГГГ-ММ.', 'всегда', None),
        ('status', 'Статус', 'Где строка плана сейчас.', '', 'list:plan_status'),
        ('checked', 'Дата проверки', 'ГГГГ-ММ-ДД или ГГГГ-ММ.', 'всегда', None),
    ]),
    ('Изменения', 'changes.csv', 'одно изменение записи при обновлении', '/process/result-and-upkeep/', [
        ('version', 'Версия', 'Версия реестра — удобно месяцем выпуска: 2027-04.', 'всегда', None),
        ('date', 'Дата', 'ГГГГ-ММ-ДД.', 'всегда', None),
        ('record_id', 'Запись', 'id записи реестра.', 'всегда', 'ref?:Реестр'),
        ('change', 'Вид изменения', 'Новая запись, изменено поле, изменён статус.', 'всегда', 'list:change'),
        ('field', 'Поле', 'Код колонки реестра.', '', None),
        ('old_value', 'Было', 'Прежнее значение.', '', None),
        ('new_value', 'Стало', 'Новое значение.', '', None),
        ('reason', 'Почему', 'Откуда известно об изменении.', 'всегда', None),
        ('author', 'Кто внёс', 'Роль или имя.', '', None),
    ]),
]

INT_COLS = {('Вопросы', 'n'), ('Связи', 'question_n'), ('Журнал поиска', 'question_n'), ('План', 'question_n'),
            ('Журнал поиска', 'viewed'), ('Проверки качества', 'question')}

FILL_REQ = PatternFill('solid', fgColor='FCE4D6')
FILL_HDR = PatternFill('solid', fgColor='E7E6E6')
FILL_CODE = PatternFill('solid', fgColor='F2F2F2')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(bottom=THIN)


def read_csv(name):
    with open(os.path.join(T, name), encoding='utf-8', newline='') as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def check_against_schema():
    schema = json.load(open(os.path.join(T, 'registry.schema.json'), encoding='utf-8'))
    names = [f['name'] for f in schema['fields']]
    assert names == [c[0] for c in REGISTRY], 'колонки реестра расходятся со схемой'
    req = {f['name'] for f in schema['fields'] if f.get('constraints', {}).get('required')}
    mine = {c[0] for c in REGISTRY if c[3]}
    assert req == mine, (req ^ mine)
    enums = {f['name']: f['constraints']['enum'] for f in schema['fields'] if 'enum' in f.get('constraints', {})}
    for code, _, _, _, ref in REGISTRY:
        if code in enums:
            lst = ref.split(':')[1]
            assert [c for c, _ in LISTS[lst][1]] == enums[code], code


def build_workbook():
    check_against_schema()
    wb = Workbook()
    ws0 = wb.active
    ws0.title = 'Как заполнять'

    # справочники
    wl = wb.create_sheet('Справочники')
    wl['A1'] = 'Справочники кодов. Выпадающие списки в листах берут значения отсюда.'
    wl['A1'].font = Font(bold=True)
    col = 1
    for name, (title, values) in LISTS.items():
        c = get_column_letter(col)
        wl.cell(row=2, column=col, value=f'{title} — {name}').font = Font(bold=True)
        wl.cell(row=2, column=col).fill = FILL_HDR
        wl.cell(row=2, column=col + 1, value='что значит').fill = FILL_HDR
        for i, (code, meaning) in enumerate(values):
            wl.cell(row=3 + i, column=col, value=code)
            wl.cell(row=3 + i, column=col + 1, value=meaning)
        rng = f"'Справочники'!${c}$3:${c}${2 + len(values)}"
        wb.defined_names[name] = DefinedName(name, attr_text=rng)
        wl.column_dimensions[c].width = 22
        wl.column_dimensions[get_column_letter(col + 1)].width = 48
        col += 3
    wl.freeze_panes = 'A3'

    sheet_objs = {}
    for title, fname, row_is, chapter, cols in SHEETS:
        ws = wb.create_sheet(title, len(wb.sheetnames) - 1)
        sheet_objs[title] = ws
        header, rows = read_csv(fname)
        assert header == [c[0] for c in cols], (fname, header)
        for j, (code, ru, hint, req, ref) in enumerate(cols, 1):
            letter = get_column_letter(j)
            label = ru + (' *' if req else '')
            h = ws.cell(row=1, column=j, value=label)
            h.font = Font(bold=True)
            h.fill = FILL_REQ if req else FILL_HDR
            h.alignment = Alignment(wrap_text=True, vertical='top')
            note = hint
            if isinstance(req, int):
                note = f'Обязательно, атрибут № {req} из 14. ' + hint
            elif req == 'всегда':
                note = 'Обязательно. ' + hint
            h.comment = Comment(note, 'Инвентаризация данных', width=320, height=140)
            k = ws.cell(row=2, column=j, value=code)
            k.font = Font(name='Consolas', size=9, color='595959')
            k.fill = FILL_CODE
            k.border = BORDER
            if (title, code) not in INT_COLS:
                ws.column_dimensions[letter].number_format = '@'
            width = max(len(ru) + 4, len(code) + 2, 12)
            ws.column_dimensions[letter].width = min(width, 28)
            # проверка значения и подсказка при выборе ячейки
            sq = f'{letter}3:{letter}{MAXROW}'
            prompt = note[:250]
            if ref and ref.startswith('list'):
                strict = ref.startswith('list:')
                lst = ref.split(':')[1]
                dv = DataValidation(type='list', formula1=f'={lst}', allow_blank=True,
                                    errorStyle='stop' if strict else 'warning',
                                    showErrorMessage=True, showInputMessage=True,
                                    promptTitle=ru[:32], prompt=prompt,
                                    errorTitle='Код не из справочника',
                                    error='Выберите код из списка: значения и их смысл — на листе «Справочники».'
                                    + ('' if strict else ' Несколько кодов пишут через «;».'))
            elif ref and ref.startswith('ref'):
                target = ref.split(':')[1]
                dv = DataValidation(type='list', formula1=f"='{target}'!$A$3:$A${MAXROW}", allow_blank=True,
                                    errorStyle='warning', showErrorMessage=True, showInputMessage=True,
                                    promptTitle=ru[:32], prompt=prompt,
                                    errorTitle='Нет такого идентификатора',
                                    error=f'Такого идентификатора нет на листе «{target}». Проверьте или добавьте строку туда.')
            else:
                dv = DataValidation(allow_blank=True, showInputMessage=True,
                                    promptTitle=ru[:32], prompt=prompt)
            dv.add(sq)
            ws.add_data_validation(dv)
        for i, r in enumerate(rows, 3):
            for j, v in enumerate(r, 1):
                code = cols[j - 1][0]
                if (title, code) in INT_COLS and v.isdigit():
                    v = int(v)
                c = ws.cell(row=i, column=j, value=v if v != '' else None)
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = 'C3' if title == 'Реестр' else 'B3'
        ws.auto_filter.ref = f'A2:{get_column_letter(len(cols))}{max(3, len(rows) + 2)}'
        # цвет — по коду, а не вместо кода
        codes = [c[0] for c in cols]
        def cf(code, mapping):
            if code not in codes:
                return
            L = get_column_letter(codes.index(code) + 1)
            for val, color in mapping.items():
                ws.conditional_formatting.add(f'{L}3:{L}{MAXROW}', CellIsRule(
                    operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=color)))
        access_colors = {'open': 'C6EFCE', 'unstructured': 'E2F0D9', 'restricted': 'FFEB9C',
                         'closed': 'FFC7CE', 'unknown': 'D9D9D9'}
        if title == 'Реестр':
            cf('access_level', access_colors)
        if title == 'Связи':
            cf('status', access_colors)
            cf('quality', {'fit': 'C6EFCE', 'fit_with_caveats': 'FFEB9C', 'unfit': 'FFC7CE', 'not_checked': 'D9D9D9'})
        if title == 'План':
            cf('status', {'done': 'C6EFCE', 'in_progress': 'FFEB9C', 'planned': 'E7E6E6', 'cancelled': 'D9D9D9'})
        if title == 'Проверки качества':
            cf('impact', {'none': 'C6EFCE', 'caveat': 'FFEB9C', 'blocking': 'FFC7CE'})

    # лист «Колонки» — все подсказки на виду
    wc = wb.create_sheet('Колонки', len(wb.sheetnames) - 1)
    hdr = ['Лист', 'Код колонки', 'Название', 'Обязательно', 'Что писать', 'Значения']
    for j, h in enumerate(hdr, 1):
        c = wc.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = FILL_HDR
    r = 2
    for title, fname, row_is, chapter, cols in SHEETS:
        for code, ru, hint, req, ref in cols:
            reqs = f'да, № {req}' if isinstance(req, int) else ('да' if req else '')
            vals = ''
            if ref and ref.startswith('list'):
                vals = ', '.join(c for c, _ in LISTS[ref.split(':')[1]][1])
            elif ref and ref.startswith('ref'):
                vals = 'id с листа «' + ref.split(':')[1] + '»'
            for j, v in enumerate([title, code, ru, reqs, hint, vals], 1):
                c = wc.cell(row=r, column=j, value=v)
                c.alignment = Alignment(wrap_text=True, vertical='top')
            r += 1
    for L, w in zip('ABCDEF', (18, 20, 26, 12, 70, 40)):
        wc.column_dimensions[L].width = w
    wc.freeze_panes = 'A2'
    wc.auto_filter.ref = f'A1:F{r - 1}'

    # первый лист — как заполнять
    lines = [
        ('Рабочая таблица инвентаризации данных', 'title'),
        (f'Шаблон из книги «Инвентаризация данных»: {BOOK}/appendices/templates/', ''),
        ('', ''),
        ('Как устроена', 'h'),
        ('Лист — одна сущность, строка — одна вещь. Листы связаны идентификаторами: колонка одного листа хранит id строки другого. Схема связей — на странице «Шаблоны».', ''),
        ('Строка 1 — название колонки, строка 2 — её код, как в CSV-шаблоне. Наведите курсор на название — появится подсказка; она же видна при выборе ячейки. Все подсказки — на листе «Колонки».', ''),
        ('Колонки со звёздочкой и оранжевой шапкой обязательны. В реестре это 14 атрибутов записи, дата проверки и откуда известно.', ''),
        ('Пустую обязательную ячейку не оставляйте: пишите значение, «нет» (сведения не существует) или «не установлено» и где искать; в колонках с кодами — unknown.', ''),
        ('Коды выбирайте из выпадающего списка. Что они значат — на листе «Справочники». Цвет ячейки ставится сам по коду; цвет без кода не пишите — при выгрузке в CSV он пропадёт.', ''),
        ('Даты — 2026-10-25 или 2026-10. Колонки с датами и кодами — текстовые, чтобы таблица их не переделала.', ''),
        ('Строки начиная с третьей — примеры из книги: реестр воздушных судов FAA и условный пример о доступности медицинской помощи. Удалите их перед работой.', ''),
        ('', ''),
        ('Как отдать лист программе', 'h'),
        ('Удалите строку 1 с русскими названиями, сохраните лист как «CSV UTF-8 (разделитель — запятая)». Получится файл с колонками, как в шаблоне CSV; реестр проверяют по схеме registry.schema.json.', ''),
        ('', ''),
        ('Листы', 'h'),
    ]
    for i, (text, kind) in enumerate(lines, 1):
        c = ws0.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'title':
            c.font = Font(bold=True, size=14)
        elif kind == 'h':
            c.font = Font(bold=True, size=12)
    r = len(lines) + 1
    for j, h in enumerate(['Лист', 'CSV-шаблон', 'Строка — это', 'Глава книги'], 1):
        c = ws0.cell(row=r, column=j, value=h); c.font = Font(bold=True); c.fill = FILL_HDR
    for title, fname, row_is, chapter, cols in SHEETS:
        r += 1
        ws0.cell(row=r, column=1, value=title)
        ws0.cell(row=r, column=2, value=fname)
        ws0.cell(row=r, column=3, value=row_is)
        link = ws0.cell(row=r, column=4, value=BOOK + chapter)
        link.hyperlink = BOOK + chapter
        link.font = Font(color='0563C1', underline='single')
    for L, w in zip('ABCD', (60, 20, 40, 50)):
        ws0.column_dimensions[L].width = w

    wb.properties.title = 'Рабочая таблица инвентаризации данных'
    wb.properties.creator = 'Инвентаризация данных — datainventory.begtin.tech'
    out = os.path.join(T, 'workbook.xlsx')
    wb.save(out)
    return out


# ---------------------------------------------------------------- Word
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def new_doc(title, subtitle=None, landscape=False):
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Arial'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for s in ('Heading 1', 'Heading 2', 'Title'):
        d.styles[s].font.name = 'Arial'
        d.styles[s].font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ('left_margin', 'right_margin'):
        setattr(sec, side, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    d.core_properties.title = title
    d.core_properties.author = 'Инвентаризация данных — datainventory.begtin.tech'
    h = d.add_heading(title, level=1)
    if subtitle:
        p = d.add_paragraph(subtitle)
        p.runs[0].italic = True
    footer = sec.footer.paragraphs[0]
    footer.text = f'Шаблон из книги «Инвентаризация данных» — {BOOK}/appendices/templates/ · CC BY 4.0'
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return d


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def table(d, rows, widths=None, header=True, fill_first_col=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        t.autofit = False
        for j, w in enumerate(widths):
            t.columns[j].width = w
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ''
            p = c.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9.5)
            if header and i == 0:
                run.bold = True
                shade(c, 'E7E6E6')
            if fill_first_col and j == 0 and not (header and i == 0):
                run.bold = True
                shade(c, 'F2F2F2')
            if widths:
                c.width = widths[j]
    d.add_paragraph()
    return t


def para(d, text, bold_prefix=None, size=None, italic=False):
    p = d.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        for run in p.runs:
            run.font.size = Pt(size)
    return p


def bullets(d, items):
    for it in items:
        if isinstance(it, tuple):
            p = d.add_paragraph(style='List Bullet')
            r = p.add_run(it[0]); r.bold = True
            p.add_run(it[1])
        else:
            d.add_paragraph(it, style='List Bullet')


def numbered(d, items):
    for it in items:
        d.add_paragraph(it, style='List Number')


BOX = '☐'


def build_questionnaire():
    d = new_doc('Анкета владельца данных',
                'Образец: условный пример книги — министерство здравоохранения, тема «доступность медицинской помощи». '
                'Замените строки своими, вписанными заранее.')
    para(d, 'Подразделение: ____________________   Кто заполняет (ФИО, должность): ____________________   Дата: ____________')
    d.add_heading('Как заполнить', level=2)
    numbered(d, [
        'Мы уже вписали наборы данных, которые нашли в документах вашего подразделения: положениях, контрактах, на сайтах. '
        'У каждой строки указано, откуда она.',
        'У каждого набора поставьте одну отметку. Строка без отметки — не ответ.',
        'Исправьте неверное прямо в карточке. Если не знаете ответа — напишите «не знаю» и, если можете, кто знает.',
        'Впишите наборы, которых у нас нет: таблицы, журналы и списки, которые вы ведёте вне систем, на общем диске, в почте, на бумаге.',
        'Ответьте на три вопроса в конце. Анкета займёт около 20 минут; по ней мы проведём короткое интервью.',
    ])
    table(d, [
        ['Отметка', 'Когда ставить'],
        [f'{BOX} Подтверждаю', 'Всё в карточке верно'],
        [f'{BOX} Исправляю', 'Верно с поправками — поправки впишите в карточку'],
        [f'{BOX} Не наше', 'Набор ведёт другое подразделение или организация — напишите, кто'],
        [f'{BOX} Не ведём', 'Набор не ведут или перестали вести — напишите, когда и куда делись данные'],
        ['Новое', 'Ставится у наборов, которые вы вписали сами'],
        ['Не хватает', 'Данные нужны вам для работы, но их нет'],
    ], widths=[Cm(4), Cm(13)])

    def card(n, title, src, values, sample=False):
        h = d.add_heading(f'Набор {n}. {title}' if title else f'Набор {n}', level=2)
        if sample:
            para(d, 'Образец заполнения', italic=True)
        rows = [['Вопрос', 'Ответ']]
        rows.append(['Откуда строка', src])
        qs = [
            '1. В какой системе ведётся набор? Если вне систем — где: таблица, документ, почта, бумага',
            '2. Как называется набор и что в нём: люди, организации, объекты, события; за какой период и территорию',
            '3. Кто ведёт набор и может ответить на вопросы о нём — ФИО, должность, контакты',
            '4. Как часто обновляются данные: ежедневно, ежемесячно, ежегодно, не обновляются, не знаю',
            '5. Откуда поступают данные: вносите сами, получаете от других организаций, покупаете',
            '6. Кому и как данные передаются: другим ведомствам, гражданам, публично; есть ли ссылка',
            '7. Есть ли в наборе сведения о людях: да, нет, не знаю',
            '8. Есть ли другие ограничения доступа — тайна, договор? На каком основании',
            '9. Есть ли описание структуры — схема, справочник полей: да, нет, не знаю',
        ]
        for q, v in zip(qs, values):
            rows.append([q, v])
        rows.append(['Отметка', f'{BOX} Подтверждаю   {BOX} Исправляю   {BOX} Не наше   {BOX} Не ведём'])
        rows.append(['Комментарий', values[9] if len(values) > 9 else ''])
        t = table(d, rows, widths=[Cm(8), Cm(9)], fill_first_col=True)
        if sample:
            # отметка в образце — «Подтверждаю»
            t.cell(len(rows) - 2, 1).paragraphs[0].runs[0].text = '☒ Подтверждаю   ☐ Исправляю   ☐ Не наше   ☐ Не ведём'

    # строки — из questionnaire.csv
    hdr, rows = read_csv('questionnaire.csv')
    R = [dict(zip(hdr, r)) for r in rows]
    freq = {'daily': 'ежедневно', 'monthly': 'ежемесячно', 'annual': 'ежегодно', 'continuous': 'непрерывно'}
    yn = {'yes': 'да', 'no': 'нет', 'unknown': 'не знаю', '': ''}
    first = R[0]
    card(1, first['dataset'], first['prefilled_from'], [
        first['system'],
        first['description'],
        'начальник ИТ-отдела центра медицинской статистики; рабочая почта и телефон',
        freq.get(first['update_frequency'], first['update_frequency']),
        first['received_from'],
        first['shared_with'],
        yn[first['personal_data']],
        f"{first['access_restrictions']}; {first['restriction_basis']}",
        yn[first['schema_available']],
        '',
    ], sample=True)
    card(2, '', 'впишите: статья положения, номер контракта или адрес страницы', [''] * 10)

    d.add_heading('Новые наборы', level=2)
    para(d, 'Таблицы, журналы и списки, которые ведёте вы, но которых нет выше. По строке на набор.')
    table(d, [['Как называется и что в нём', 'Где хранится', 'Кто ведёт', 'Как часто обновляется', 'Сведения о людях']]
          + [['', '', '', '', ''] for _ in range(4)], widths=[Cm(6), Cm(3), Cm(3), Cm(2.5), Cm(2.5)])
    para(d, 'Образец из условного примера: «Журнал ожидания плановой госпитализации — пациенты, направленные на плановую '
            'госпитализацию: дата направления, дата госпитализации, медицинская организация» · таблица на общем диске '
            'отдела · заместитель начальника отдела · ежемесячно · да.', italic=True, size=9)

    d.add_heading('Три вопроса', level=2)
    for q in ['Какие таблицы, журналы и списки вы ведёте вне систем? (впишите их в «Новые наборы»)',
              'Каких данных вам не хватает для работы? Образец: «сведения о пациентах из региона, которые лечатся в других регионах».',
              'Кто ещё работает с этими данными — в вашей организации и за её пределами?']:
        para(d, q)
        para(d, '_' * 95)
    out = os.path.join(T, 'questionnaire.docx')
    d.save(out)
    return out


def build_cover_letter():
    d = new_doc('Сопроводительное письмо к анкете',
                'Образец на 8–10 строк. Отправляет и подписывает владелец задачи, а не проектная команда. '
                'Квадратные скобки — заменить своим.')
    para(d, 'Тема: Инвентаризация данных [организации] — анкета для [подразделения], ответ до [дата]')
    para(d, 'Уважаемый(ая) [имя отчество],')
    for t in [
        'по поручению [кто поручил, документ] мы проводим инвентаризацию данных: выясняем, какие данные есть в [организации], '
        'кто их ведёт и на каком основании. Это нужно, чтобы [цель одной фразой — например, найти данные для оценки доступности '
        'медицинской помощи в регионах].',
        'Во вложении — анкета. Мы уже вписали в неё наборы, которые нашли в документах вашего подразделения. Подтвердите или '
        'исправьте каждую строку и добавьте то, что ведёте вне систем. Если ответа нет — напишите «не знаю». Заполнение займёт '
        'около 20 минут.',
        'Пожалуйста, ответьте до [дата]. Интервью по анкете — 15–30 минут — предлагаем провести [дата и время]; если неудобно, '
        'предложите другое время.',
        'Ответы увидят [кто]. В общедоступную часть реестра попадут [что именно — например, название, описание и режим доступа '
        'наборов; контакты сотрудников и внутренние адреса — нет].',
        'Вопросы — [имя, телефон, почта].',
    ]:
        para(d, t)
    para(d, '[Подпись владельца задачи, должность]')
    d.add_heading('Перед отправкой', level=2)
    bullets(d, [
        'Срок — дата, а не «до конца квартала»; не в отпуска и не в отчётные недели.',
        'Напоминание — до срока; о тех, кто не ответил к сроку, узнаёт владелец задачи.',
        'Сначала одна-две очереди: ошибку в анкете дешевле найти на одном подразделении.',
        'Тем, у кого данных нет, тоже нужно ответить: ответ «данных не ведём» записывают.',
    ])
    out = os.path.join(T, 'cover-letter.docx'); d.save(out); return out


def build_access_request():
    d = new_doc('Запрос владельцу данных',
                'Образец запроса к чужому владельцу — государственному органу или организации. '
                'Квадратные скобки — заменить своим.')
    para(d, 'Кому: [орган или организация, должность, имя]')
    para(d, 'От: [имя или организация, адрес для ответа]')
    para(d, 'Тема: Запрос о наборе данных «[название]»')
    para(d, 'Прошу на основании [закон о доступе к информации страны, статья] сообщить сведения о наборе данных:')
    numbered(d, [
        'Ведёт ли [орган] сведения о [что: объекты, события, показатели] за [период] по [территория]? В какой системе или каком документе?',
        'Опубликованы ли они? Если да — по какому адресу и в каком формате.',
        'Если не опубликованы — можно ли их получить, в каком виде (весь набор, агрегаты, выборка) и на каких условиях: запрос, соглашение, плата, срок.',
        'Если доступ ограничен — на основании какой нормы (закон, статья) и какие сведения в наборе она закрывает.',
        'Кто отвечает за эти данные и может ответить на вопросы о них?',
    ])
    para(d, 'Сведения нужны для [зачем — одной фразой]. Ответ прошу направить [способ] в срок, установленный законом.')
    para(d, '[Дата, подпись]')
    d.add_heading('Как пользоваться', level=2)
    bullets(d, [
        ('Три–пять вопросов о конкретном наборе. ', 'Закон даёт право на информацию и документы, а не на опись: на вопрос «какие у вас есть данные?» обычно отвечают ссылкой на сайт.'),
        ('Сначала поищите. ', 'Перечни информации и документов органа, ответы на чужие запросы на общественных сайтах — готовое основание режима доступа.'),
        ('Закон и сроки ответа ', f'по странам — в приложении «Нормативная база»: {BOOK}/appendices/legal-framework/'),
        ('Ответ записывают с датой. ', '«Данных нет» — пробел с основанием. Нет ответа — в основании дата запроса и «ответа нет», режим доступа остаётся «не установлено».'),
        ('Компании отвечать не обязаны. ', 'Поставщику данных пишут как покупатель — по шаблону «Вопросы поставщику данных».'),
    ])
    out = os.path.join(T, 'access-request.docx'); d.save(out); return out


def build_interview_guide():
    d = new_doc('Памятка для интервью',
                'Интервью — не второй лист, а разговор по заполненной анкете: 15–30 минут. Распечатайте анкету адресата заранее.')
    d.add_heading('Перед встречей', level=2)
    bullets(d, [
        'Интервью назначают при рассылке анкеты, а не после ответа.',
        'Вдвоём: один спрашивает, другой сразу вписывает ответы в анкету.',
        'Ведёт свой сотрудник — он знает людей и предмет; подрядчик — вместе с ним.',
        'Отметьте в анкете строки для разговора: новые, вычеркнутые, исправленные, «не знаю» и пустые ячейки.',
    ])
    d.add_heading('План на 30 минут', level=2)
    table(d, [
        ['Минуты', 'Что делать'],
        ['0–3', 'Зачем инвентаризация, кто поручил, что попадёт в общедоступную часть реестра. Запись разговора — только с согласия.'],
        ['3–15', 'По строкам анкеты: сначала новые — главная находка опроса, — потом вычеркнутые, исправленные и «не знаю». Просите показать: систему, общий диск, таблицу, из которой собирают отчёт.'],
        ['15–25', 'О работе, а не о «данных»: что подразделение отправляет каждый месяц и кому, что получает, по каким таблицам отвечает на запросы руководства. Каких данных не хватает. Кто ещё работает с этими данными.'],
        ['25–30', 'Что изменилось в анкете, кто подтвердит строки письменно и к какой дате. Кому ещё задать вопросы.'],
    ], widths=[Cm(2.5), Cm(14.5)])
    para(d, 'Если на всё 15 минут — только новые и вычеркнутые строки и вопрос «каких данных не хватает».', italic=True)
    d.add_heading('Что спросить', level=2)
    bullets(d, [
        'Какой отчёт вы готовите каждый месяц и из чего?',
        'Какие списки и журналы ведёте? Что лежит на общем диске, в почте, в шкафу?',
        'Откуда приходят эти сведения и кому вы их передаёте?',
        'Чего вам не хватает для работы?',
        'Кто ещё работает с этими данными?',
    ])
    d.add_heading('После интервью', level=2)
    bullets(d, [
        'В тот же день отправьте ответственному исправленные строки — он подтверждает их письменно. Такой ответ — основание для записи реестра.',
        'Новые строки сверьте с выборкой, документом или экраном; до проверки — с пометкой «со слов владельца».',
        'В листе «Адресаты» — дата интервью и итог; новые адресаты — отдельными строками.',
        'Запись разговора — служебная информация: внешнему сервису, в том числе языковой модели, её не отдают без решения владельца.',
    ])
    out = os.path.join(T, 'interview-guide.docx'); d.save(out); return out


def build_goal_card():
    d = new_doc('Цель и границы инвентаризации',
                'Карточка на одну страницу. Заполняют до начала поиска и согласуют с владельцем задачи. '
                'Серым курсивом — условный пример книги; замените своим.')
    rows = [
        ['Что записать', 'Ваш ответ', 'Пример'],
        ['Цель: что сделают с данными', '', 'Найти источники для ответа на 15 вопросов руководства о доступности медицинской помощи в регионах и показать, каких данных нет'],
        ['Владелец задачи — кто примет результат', '', 'Заместитель министра здравоохранения'],
        ['Кто пользуется результатом', '', 'Департамент, который готовит справки для руководства'],
        ['Как проверить результат', '', 'У каждого из 15 вопросов — источник с доступностью и качеством или записанный пробел'],
        ['Вид результата', '', 'Карта данных: таблица «вопрос — источник — статус» и список пробелов'],
        ['Подход', '', '«Где найти нужное»'],
        ['Организации и уровни', '', 'Государственные источники центрального и регионального уровня'],
        ['Территория и период', '', 'Вся страна по регионам; ряды за последние десять лет'],
        ['Слои данных', '', 'Государственные; международные — для сравнения с другими странами; коммерческих нет'],
        ['Глубина описания', '', 'До наборов'],
        ['Что исключено и почему', '', 'Частная медицина — следующей очередью; лекарственное обеспечение как отдельная тема'],
        ['Кто согласовал и когда', '', 'Заместитель министра, 29 сентября 2026 года'],
    ]
    t = table(d, rows, widths=[Cm(5), Cm(6), Cm(6)], fill_first_col=True)
    for i in range(1, len(rows)):
        run = t.cell(i, 2).paragraphs[0].runs[0]
        run.italic = True
        run.font.color.rgb = RGBColor(0x76, 0x76, 0x76)
    para(d, 'Строку «что исключено» забывают чаще остальных. Без неё через год не отличить пропуск от решения.', italic=True)
    d.add_heading('Вопросы к данным', level=2)
    para(d, 'Ведут листом «Вопросы» рабочей таблицы: постоянный номер, текст в укрупнённой форме, типовая тема, задача, '
            'приоритет и порог «годится» — до проверки.')
    table(d, [['№', 'Вопрос', 'Типовая тема', 'Главный или попутный', 'Порог «годится»'],
              ['2', 'Где найти сведения о числе врачей по специальностям, регионам и годам — отдельно для городов и сёл?',
               'люди', 'главный', 'все регионы, ряд за пять лет, разрез «город — село»'],
              ['4', 'Где найти сведения о сроках ожидания медицинской помощи по регионам?', 'статистика', 'главный',
               'дата направления и дата помощи заполнены не меньше чем у 90% записей; данные не старше квартала'],
              ['', '', '', '', ''], ['', '', '', '', '']],
          widths=[Cm(1), Cm(6), Cm(2.5), Cm(2.5), Cm(5)])
    out = os.path.join(T, 'goal-and-scope.docx'); d.save(out); return out


def build_vendor_questions():
    d = new_doc('Вопросы поставщику данных',
                'Для коммерческих данных: компания отвечать не обязана, поэтому спрашивают как покупатель. '
                'Каждый ответ ложится в поле реестра.', landscape=True)
    para(d, 'Поставщик: ____________________   Набор или продукт: ____________________   Дата разговора: ____________')
    rows = [['Вопрос', 'Зачем', 'Поле реестра', 'Ответ']]
    for q, why, field in [
        ('Кого покрывают данные? Доля рынка, регионы, какие клиенты', 'Выводы о стране по клиентам одной компании — главная ошибка', 'Описание; отметка о качестве'),
        ('С какого периода и с какими разрывами?', 'Изменения выборки ломают ряды', 'Описание; отметка о качестве'),
        ('Какая подробность? Территория, период, срезы, наименьший размер ячейки', 'Малые значения обычно скрыты, и маленькие территории выпадают', 'Описание; условия доступа'),
        ('Как часто и с какой задержкой?', 'Ради свежести коммерческие данные и берут', 'Периодичность'),
        ('Какая методика и как она менялась? Пересматривают ли историю', 'Пересмотр меняет уже полученные числа', 'Схема данных; отметка о качестве'),
        ('Свои это данные или перепродажа?', 'У перепродавца первоисточник другой', 'Связанные записи (copy_of:)'),
        ('Как данные обезличены?', 'От этого зависят хранение, доступ внутри организации и публикация', 'Персональные данные'),
        ('На каких условиях? Цель, срок, передача, публикация производных, что делать с данными после окончания договора', 'Нарушить договор проще всего тем, о чём не спросили', 'Условия использования; условия доступа'),
        ('Сколько стоит и сколько ждать?', 'Цена и срок решают, стоит ли добиваться данных', 'Условия доступа'),
        ('Можно ли получить образец?', 'Пригодность проверяют на данных, а не на описании', 'Отметка о качестве'),
    ]:
        rows.append([q, why, field, ''])
    table(d, rows, widths=[Cm(6.5), Cm(6.5), Cm(4), Cm(8)])
    d.add_heading('После разговора', level=2)
    bullets(d, [
        'Запишите номер и дату соглашения или ссылку на условия — и для бесплатного источника: условия меняются.',
        'Срок подписки и что делать с данными после него: удалить, хранить без обновления, продлить.',
        'Кто в организации может работать с данными, что можно публиковать и кому передавать.',
        'Образец проверьте как любые данные: выборка, разрывы в рядах, выбросы и пустые значения, пересмотры методики.',
    ])
    out = os.path.join(T, 'vendor-questions.docx'); d.save(out); return out


ZIP_NAME = 'data-inventory-templates.zip'


def build_zip():
    files = sorted(f for f in os.listdir(T) if not f.startswith('.') and f != ZIP_NAME)
    out = os.path.join(T, ZIP_NAME)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(os.path.join(T, f), arcname=f'data-inventory-templates/{f}')
    return out, files


if __name__ == '__main__':
    for fn in (build_workbook, build_questionnaire, build_cover_letter, build_access_request,
               build_interview_guide, build_goal_card, build_vendor_questions):
        print(fn())
    out, files = build_zip()
    print(out, files)
